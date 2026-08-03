from __future__ import annotations

import csv
import io
from typing import Any

from fastapi import HTTPException, status
from openpyxl import load_workbook
from pydantic import ValidationError

from core.menus.schemas import DishCreate

REQUIRED_COLUMNS = ("meal_code", "name_es", "name_en")
OPTIONAL_COLUMNS = (
    "description_es",
    "description_en",
    "image_url",
    "availability",
)
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def parse_dish_upload(
    content: bytes, filename: str
) -> tuple[list[DishCreate], list[dict[str, Any]]]:
    """Parse CSV or Excel into validated dishes and per-row errors."""
    name = filename.lower()
    if name.endswith(".csv"):
        rows = _read_csv_rows(content)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        rows = _read_excel_rows(content)
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Use .csv or .xlsx",
        )

    if not rows:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="File has no data rows",
        )

    dishes: list[DishCreate] = []
    errors: list[dict[str, Any]] = []

    for index, raw in enumerate(rows, start=2):
        if _row_is_empty(raw):
            continue
        missing = [col for col in REQUIRED_COLUMNS if not str(raw.get(col) or "").strip()]
        if missing:
            errors.append(
                {
                    "row": index,
                    "detail": f"missing required columns: {', '.join(missing)}",
                }
            )
            continue
        try:
            dishes.append(DishCreate.model_validate(raw))
        except ValidationError as exc:
            errors.append({"row": index, "detail": _first_validation_message(exc)})

    if not dishes and errors:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail={
                "message": "No valid dishes in file",
                "errors": errors,
            },
        )
    if not dishes:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="File has no data rows",
        )

    return dishes, errors


def _read_csv_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="CSV must be UTF-8 encoded",
        ) from exc

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="CSV header row is missing",
        )
    _ensure_headers(reader.fieldnames)
    return [dict(row) for row in reader]


def _read_excel_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Invalid Excel file",
        ) from exc

    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Excel sheet is empty",
        ) from exc

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    _ensure_headers(headers)

    parsed: list[dict[str, Any]] = []
    for values in rows_iter:
        row: dict[str, Any] = {}
        for idx, key in enumerate(headers):
            if key not in ALL_COLUMNS:
                continue
            value = values[idx] if idx < len(values) else None
            if value is None:
                row[key] = None
            else:
                row[key] = str(value).strip() if not isinstance(value, str) else value
        parsed.append(row)
    return parsed


def _ensure_headers(fieldnames: list[str] | None) -> None:
    if not fieldnames:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Header row is missing",
        )
    normalized = [name.strip() for name in fieldnames if name and name.strip()]
    missing = [col for col in REQUIRED_COLUMNS if col not in normalized]
    if missing:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"Missing required columns: {', '.join(missing)}",
        )


def _row_is_empty(row: dict[str, Any]) -> bool:
    return not any(str(value or "").strip() for value in row.values())


def _first_validation_message(exc: ValidationError) -> str:
    errors = exc.errors()
    if not errors:
        return "invalid row"
    err = errors[0]
    loc = ".".join(str(part) for part in err.get("loc", ()))
    msg = err.get("msg", "invalid value")
    return f"{loc}: {msg}" if loc else msg
